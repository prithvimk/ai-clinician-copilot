"""Generate a synthetic local care pathway (DOCX) to exercise DOCX ingestion.

Content paraphrases the KDIGO 2022 diabetes-in-CKD and WHO 2021 hypertension guidelines in
the corpus; it is a software demo fixture, not clinical guidance.
"""

from pathlib import Path

import docx

OUT = Path(__file__).resolve().parents[1] / "data" / "guidelines" / "sample-care-pathway.docx"

SECTIONS = {
    "Scope": [
        "This pathway covers adults in primary care with type 2 diabetes, hypertension, "
        "or both, including people with chronic kidney disease (CKD).",
    ],
    "Type 2 Diabetes with CKD": {
        "Glycaemic Targets and Monitoring": [
            "Individualise the HbA1c target from below 6.5% to below 8.0% based on "
            "hypoglycaemia risk, comorbidities, life expectancy and patient preference.",
            "Measure HbA1c twice a year when at target and stable; measure up to four times a "
            "year when not at target or after a change in glucose-lowering therapy.",
        ],
        "Metformin and Kidney Function": [
            "Use metformin as first-line therapy when eGFR is 30 ml/min/1.73m2 or above.",
            "Reduce the metformin dose when eGFR is 30-44 ml/min/1.73m2, and stop metformin "
            "when eGFR falls below 30 ml/min/1.73m2 or the patient is on dialysis.",
            "Monitor eGFR at least annually on metformin, and every 3-6 months when eGFR is "
            "below 60 ml/min/1.73m2.",
        ],
        "SGLT2 Inhibitors": [
            "Start an SGLT2 inhibitor for kidney and cardiovascular protection when eGFR is "
            "20 ml/min/1.73m2 or above, irrespective of HbA1c.",
        ],
    },
    "Hypertension": {
        "Treatment Targets": [
            "Target blood pressure below 140/90 mmHg for all adults with hypertension without "
            "comorbidities.",
            "Target systolic blood pressure below 130 mmHg in patients with known "
            "cardiovascular disease, diabetes, CKD or high cardiovascular risk.",
        ],
        "Follow-up and Monitoring": [
            "Review monthly after starting or changing antihypertensive medication until "
            "the target is reached.",
            "Once at target, review every 3-6 months.",
        ],
    },
}


def add_section(document: docx.Document, heading: str, body, level: int) -> None:
    document.add_heading(heading, level=level)
    if isinstance(body, dict):
        for sub_heading, sub_body in body.items():
            add_section(document, sub_heading, sub_body, level + 1)
    else:
        for paragraph in body:
            document.add_paragraph(paragraph)


def main() -> None:
    document = docx.Document()
    document.core_properties.title = "Sample Primary Care Pathway: Diabetes, CKD and Hypertension"
    document.add_heading(document.core_properties.title, level=0)
    document.add_paragraph("SYNTHETIC SAMPLE — software demonstration only, not clinical advice.")
    for heading, body in SECTIONS.items():
        add_section(document, heading, body, level=1)

    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = (
        "eGFR (ml/min/1.73m2)",
        "Metformin",
        "SGLT2 inhibitor",
    )
    for row in (
        ("≥45", "Full dose", "Start"),
        ("30-44", "Reduce dose", "Start"),
        ("20-29", "Stop", "Start"),
        ("<20", "Stop", "Do not start"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            cell.text = value

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
