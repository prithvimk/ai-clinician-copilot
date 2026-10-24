from pathlib import Path

import docx
import pymupdf
import pytest

BODY = (
    "Patients with uncontrolled type 2 diabetes should have therapy intensified. "
    "Measure HbA1c every three months until the target is reached. "
)


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """Three-page PDF: a cover, then headed sections in 16pt over 10pt body text."""
    pdf = pymupdf.open()
    cover = pdf.new_page()
    cover.insert_text((72, 200), "Diabetes Guideline", fontsize=28)

    for page_no, (heading, body) in enumerate(
        [("1 Glycaemic control", BODY * 4), ("2 Blood pressure", "Target below 140/90 mmHg. " * 8)],
        start=2,
    ):
        page = pdf.new_page()
        page.insert_text((72, 100), heading, fontsize=16)
        y = 130
        for i in range(0, len(body), 80):
            page.insert_text((72, y), body[i : i + 80], fontsize=10)
            y += 13
        page.insert_text((72, page.rect.height - 20), str(page_no), fontsize=9)

    path = tmp_path / "guideline.pdf"
    pdf.save(path)
    return path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Care Pathway", level=0)
    document.add_heading("Diabetes", level=1)
    document.add_heading("Monitoring", level=2)
    document.add_paragraph(BODY)
    document.add_heading("Hypertension", level=1)
    document.add_paragraph("Review monthly until blood pressure is at target.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "eGFR", "Metformin"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "<30", "Stop"
    path = tmp_path / "pathway.docx"
    document.save(path)
    return path
